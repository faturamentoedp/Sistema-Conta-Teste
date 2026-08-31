"""
============================================================================
 GERADOR DA BASE TARIFARIA  ->  src/lib/sp/tarifas-geradas.ts
============================================================================

 Le a planilha "Tabela Tarifas Consolidada ES e SP.xlsx" (aba Export) e gera
 o arquivo TypeScript com todo o historico de tarifas por vigencia.

 COMO USAR quando sair uma resolucao nova:
   1. Substitua a planilha em Planilhas/ pela versao atualizada.
   2. Rode:   python scripts/gerar_tarifas.py
   3. Pronto. Nao precisa editar TypeScript na mao.

 A planilha NAO traz PIS/COFINS. Essas aliquotas continuam sendo mantidas
 a mao em src/lib/sp/dados.ts (TRIBUTOS_MENSAIS).
============================================================================
"""
import openpyxl, os, sys, datetime

sys.stdout.reconfigure(encoding='utf-8')

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORIGEM = os.path.join(RAIZ, 'Planilhas', 'Tabela Tarifas Consolidada ES e SP.xlsx')
DESTINO = os.path.join(RAIZ, 'src', 'lib', 'sp', 'tarifas-geradas.ts')

# A planilha usa "B2RUIRRGTB"; o resto do sistema usa "B2RUIRRG_TB".
NORMALIZAR_CATEGORIA = {
    'B2RUIRRGTB': 'B2RUIRRG_TB',
    'B1BRN/B1BPC/B1BRQ/B1BRI': 'B1BR',
}

# Subclasse -> sufixo de faixa. Categorias com faixas viram codigos distintos.
FAIXA_POR_SUBCLASSE = {
    'Residencial Desconto Social – faixa 01': 'F1',
    'Residencial Desconto Social – faixa 02': 'F2',
    'Residencial Tarifa Social – faixa 01': 'F1',
    'Residencial Tarifa Social – faixa 02': 'F2',
    'Baixa Renda': 'BR',
}

SIGLA_PARA_UF = {'EDP SP': 'SP', 'EDP ES': 'ES'}


def data_iso(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    return str(v)[:10]


def main():
    if not os.path.exists(ORIGEM):
        sys.exit(f"ERRO: planilha nao encontrada em {ORIGEM}")

    ws = openpyxl.load_workbook(ORIGEM, data_only=True)['Export']
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    registros = []
    resolucoes = {}
    for r in range(2, ws.max_row + 1):
        d = {hdr[c - 1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
        uf = SIGLA_PARA_UF.get(d.get('Sigla'))
        if not uf:
            continue

        cat = str(d['Tarifa']).strip()
        cat = NORMALIZAR_CATEGORIA.get(cat, cat)
        subclasse = str(d['Subclasse']).strip()
        faixa = FAIXA_POR_SUBCLASSE.get(subclasse, '')
        if faixa:
            cat = f'{cat}_{faixa}'

        # A planilha traz R$/MWh; o motor trabalha em R$/kWh.
        divisor = 1000.0 if str(d['Unidade']).strip().upper() == 'R$/MWH' else 1.0
        inicio = data_iso(d['Início Vigência'])
        fim = data_iso(d['Fim Vigência'])

        registros.append((
            uf, inicio, fim, cat, str(d['Posto']).strip(),
            round((d['TUSD'] or 0) / divisor, 8),
            round((d['TE'] or 0) / divisor, 8),
        ))
        resolucoes[(uf, inicio)] = str(d['Resolução ANEEL']).strip()

    # Remove duplicatas exatas (a planilha repete linhas de faixa 01/02 iguais).
    registros = sorted(set(registros), key=lambda x: (x[0], x[1], x[3], x[4]))

    ufs = sorted({x[0] for x in registros})
    print(f"Lidos {len(registros)} registros unicos ({', '.join(ufs)})")
    for uf in ufs:
        vigs = sorted({x[1] for x in registros if x[0] == uf})
        print(f"  {uf}: {len(vigs)} vigencias, de {vigs[0]} ate {vigs[-1]}")

    linhas = [
        '/**',
        ' * ARQUIVO GERADO AUTOMATICAMENTE — NAO EDITE A MAO.',
        ' *',
        ' * Origem: Planilhas/Tabela Tarifas Consolidada ES e SP.xlsx (aba Export)',
        f' * Gerado em: {datetime.date.today().isoformat()}',
        ' *',
        ' * Para atualizar: troque a planilha e rode  python scripts/gerar_tarifas.py',
        ' */',
        '',
        "export type Posto = 'Não se aplica' | 'Ponta' | 'Intermediário' | 'Fora ponta';",
        '',
        '/** [uf, inicioVigencia, fimVigencia, categoria, posto, tusd, te] */',
        'export type LinhaTarifa = readonly [',
        '  string, string, string, string, Posto, number, number,',
        '];',
        '',
        'export const TARIFAS_HISTORICO: readonly LinhaTarifa[] = [',
    ]
    for uf, ini, fim, cat, posto, tusd, te in registros:
        linhas.append(f"  ['{uf}', '{ini}', '{fim}', '{cat}', '{posto}', {tusd}, {te}],")
    linhas.append('] as const;')
    linhas.append('')
    linhas.append('/** Resolucao ANEEL por (uf, inicio de vigencia). */')
    linhas.append('export const RESOLUCOES: Readonly<Record<string, string>> = {')
    for (uf, ini), res in sorted(resolucoes.items()):
        linhas.append(f"  '{uf}|{ini}': {res!r},".replace("\\'", "'"))
    linhas.append('};')
    linhas.append('')

    os.makedirs(os.path.dirname(DESTINO), exist_ok=True)
    with open(DESTINO, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(linhas))
    print(f"\nGravado: {DESTINO}")


if __name__ == '__main__':
    main()
